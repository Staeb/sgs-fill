import base64
import io

import openpyxl
import pytest

from tests.chrome import run_js
from tests.conftest import needs_chrome

MODS = ("text", "xlsx")


def workbook_b64(build) -> str:
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def read(b64: str):
    return run_js(
        "const u8 = Uint8Array.from(atob(%r), c => c.charCodeAt(0));"
        "return await SF.readXlsx(u8.buffer);" % b64,
        modules=MODS,
    )


@needs_chrome
def test_reads_sheet_names_text_codes_and_numbers():
    def build(wb):
        ws = wb.active
        ws.title = "คะแนน"
        ws.append(["รหัส", "ก่อนกลาง", "หมายเหตุ"])
        ws.append(["07112", 18, "ปกติ"])
        ws.append([7113, 19.5, None])
        ws["A3"].number_format = "00000"
        ws.append([None, None, None])
        wb.create_sheet("ห้อง 2").append(["x"])

    out = read(workbook_b64(build))
    assert [s["name"] for s in out["sheets"]] == ["คะแนน", "ห้อง 2"]
    rows = out["sheets"][0]["rows"]
    assert rows[0] == ["รหัส", "ก่อนกลาง", "หมายเหตุ"]
    assert rows[1] == ["07112", "18", "ปกติ"]
    # Review Focus 1: รหัสที่เก็บเป็นตัวเลขจะเหลือ "7113" — Task 5 เติมศูนย์ให้
    assert rows[2] == ["7113", "19.5", ""]
    assert len(rows) == 3, "ตัดแถวว่างท้าย"


@needs_chrome
def test_rows_are_rectangular_and_sparse_cells_are_blank():
    def build(wb):
        ws = wb.active
        ws["A1"] = "a"
        ws["C1"] = "c"
        ws["B3"] = "b"

    rows = read(workbook_b64(build))["sheets"][0]["rows"]
    assert rows == [["a", "", "c"], ["", "", ""], ["", "b", ""]]


@needs_chrome
def test_thai_text_and_shared_strings_are_decoded():
    def build(wb):
        ws = wb.active
        for i in range(3):
            ws.append(["ขส", "เด็กชายสมมุติ", "ร"])

    assert read(workbook_b64(build))["sheets"][0]["rows"][2] == ["ขส", "เด็กชายสมมุติ", "ร"]


@needs_chrome
def test_a_file_that_is_not_xlsx_is_refused_in_thai():
    with pytest.raises(RuntimeError, match="ไม่ใช่ไฟล์ .xlsx"):
        run_js("await SF.readXlsx(new Uint8Array([1,2,3,4,5]).buffer);", modules=MODS)


# ── Final review, Important 3: ค่าลอยตัวจากสูตร Excel ต้องไม่หลุดเข้า SGS ──

@needs_chrome
def test_float_noise_from_formulas_is_cleaned_but_real_decimals_are_kept():
    def build(wb):
        ws = wb.active
        ws.append(["x", "y", "z", "w"])
        ws.append([0.1 + 0.2, 5.3999999999999995, 13.125, 20])

    rows = read(workbook_b64(build))["sheets"][0]["rows"]
    assert rows[1] == ["0.3", "5.4", "13.125", "20"]


def workbook_with_string_formula_b64() -> str:
    """openpyxl เขียนค่าที่คำนวณแล้วของสูตรไม่ได้ จึงแก้ XML เอง: เซลล์สูตรที่ผลเป็นข้อความ (t="str")
    เหมือนที่ Excel เก็บเมื่อสูตรดึงรหัสนักเรียนจากอีกชีต เช่น =ชื่อผู้เรียน!B29 ได้ "05923" """
    import re
    import zipfile

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "คะแนน"
    ws.append(["รหัส", "คะแนน"])
    ws.append(["=Z1", 18])
    ws.append(["=Z2", 20])
    ws["Z1"], ws["Z2"] = "x", "x"
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                x = data.decode("utf-8")
                x = re.sub(r'<c r="A2"[^>]*>.*?</c>', '<c r="A2" t="str"><f>Z1</f><v>05923</v></c>', x)
                x = re.sub(r'<c r="A3"[^>]*>.*?</c>', '<c r="A3" t="str"><f>Z2</f><v>06186</v></c>', x)
                data = x.encode("utf-8")
            zo.writestr(item, data)
    return base64.b64encode(out.getvalue()).decode()


@needs_chrome
def test_a_formula_that_returns_a_code_as_text_keeps_its_leading_zero():
    """ปพ.5 ของโรงเรียนดึงรหัสจากชีตรายชื่อด้วยสูตร ผลเป็นข้อความ "05923" — ต้องไม่กลายเป็น 5923
    ไม่งั้นรหัสเดียวกันที่ชีตอื่นเขียนเป็น "05923" จะถูกนับเป็นคนละรหัส"""
    rows = read(workbook_with_string_formula_b64())["sheets"][0]["rows"]
    assert [r[0] for r in rows[1:3]] == ["05923", "06186"]
